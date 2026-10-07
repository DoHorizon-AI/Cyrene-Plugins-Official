"""
┌─────────────────────────────────────────────────────────────────────┐
│  📄 test_office_parsing.py                                          │
│  Module: tests.test_office_parsing                                  │
│  Role: Verify real PPTX/XLSX extraction and unsupported warnings.   │
│                                                                     │
│  模块职责：验证演示文稿与工作簿的真实内容、定位和显式降级诊断。          │
└─────────────────────────────────────────────────────────────────────┘
"""

from __future__ import annotations

import base64
import hashlib
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from office_parsing import parse_office
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from parsing_types import ParseContext
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

PPTX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def test_pptx_preserves_notes_table_geometry_and_reports_visual_content(
    tmp_path: Path,
) -> None:
    """Read real presentation structures and locate all unsupported visuals."""

    source = tmp_path / "office-sample.pptx"
    _write_pptx(source)
    _inject_pptx_timing(source)

    result = parse_office(source, _context(source, PPTX_MEDIA_TYPE))

    assert result.source_format == "PPTX"
    assert result.slide_count == 1
    assert result.status == "partial_success"
    slide = result.payload["slides"][0]
    assert slide["ordinal"] == 1
    assert slide["name"] == "Executive Summary"
    assert slide["title"] == "Quarterly Overview"
    assert "speaker notes 北区" in slide["notes"]
    assert slide["slideSizeEMU"]["unit"] == "EMU"
    assert "Revenue" in result.blocks[0]["text"]
    assert "speaker notes 北区" in result.blocks[0]["text"]

    shapes = slide["shapes"]
    title_shape = next(
        shape for shape in shapes if shape.get("text") == "Quarterly Overview"
    )
    assert title_shape["shapeId"] > 0
    assert title_shape["positionEMU"]["left"] >= 0
    empty_shape = next(
        shape for shape in shapes if shape["name"] == "Empty semantic shape"
    )
    assert empty_shape["contentState"] == "empty"
    table_shape = next(shape for shape in shapes if "table" in shape)
    assert table_shape["table"]["rows"][1]["cells"][0]["text"] == "北区"

    kinds = {entry["kind"] for entry in result.unsupported_content}
    assert {"picture", "chart", "animation", "slide_transition"} <= kinds
    assert all(
        entry["locator"].get("slideOrdinal") == 1
        for entry in result.unsupported_content
    )
    assert any(diagnostic.locator for diagnostic in result.diagnostics)


def test_xlsx_keeps_sheet_order_tables_visibility_and_cached_formula_values(
    tmp_path: Path,
) -> None:
    """Preserve formula source and stored cache without calculating formulas."""

    source = tmp_path / "office-sample.xlsx"
    _write_xlsx(source)
    _set_cached_formula_values(source, {"B2": "41"}, remove={"B3"})

    result = parse_office(source, _context(source, XLSX_MEDIA_TYPE))

    assert result.source_format == "XLSX"
    assert result.sheet_count == 2
    assert result.status == "partial_success"
    assert result.payload["workbook"]["sheetOrder"] == ["Data", "Hidden"]
    sheet = result.payload["sheets"][0]
    assert sheet["ordinal"] == 1
    assert sheet["state"] == "visible"
    assert sheet["tables"][0]["name"] == "ValuesTable"
    assert sheet["tables"][0]["ref"] == "A1:B3"
    assert sheet["mergedRanges"] == ["D1:E1"]
    assert sheet["hiddenRows"] == [{"row": 4, "hidden": True}]
    assert sheet["hiddenColumns"][0]["key"] == "F"
    assert sheet["freezePanes"] == "A2"
    assert sheet["showGridLines"] is False

    formula_cells = {cell["coordinate"]: cell for cell in sheet["cells"]}
    assert formula_cells["B2"]["formula"] == "=A2*3"
    assert formula_cells["B2"]["cachedValue"] == 41
    assert formula_cells["B2"]["storedDataType"] == "f"
    assert formula_cells["B3"]["formula"] == "=SUM(A2:A3)"
    assert formula_cells["B3"]["cachedValue"] is None
    assert "formula =A2*3; cached value 41" in result.blocks[0]["text"]
    assert "formula =SUM(A2:A3); cached value unavailable" in result.blocks[0]["text"]

    missing_cache = [
        diagnostic
        for diagnostic in result.diagnostics
        if diagnostic.code == "XLSX_FORMULA_CACHE_MISSING"
    ]
    assert len(missing_cache) == 1
    assert missing_cache[0].locator["cell"] == "B3"
    assert result.payload["sheets"][1]["state"] == "hidden"


def test_xlsx_payload_and_blocks_are_utf8_json_serializable(tmp_path: Path) -> None:
    """Keep Unicode worksheet text intact in both public representations."""

    source = tmp_path / "utf8-sample.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "客户说明：已审核"
    workbook.save(source)

    result = parse_office(source, _context(source, XLSX_MEDIA_TYPE))

    assert "客户说明：已审核" in result.blocks[0]["text"]
    assert "客户说明：已审核" in str(result.payload["sheets"][0]["cells"])


def _context(source: Path, media_type: str) -> ParseContext:
    """Build an immutable parser context from the generated local fixture."""

    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    return ParseContext(
        source_ref="source:office-fixture",
        source_digest=f"sha256:{digest}",
        filename=source.name,
        media_type=media_type,
    )


def _write_pptx(path: Path) -> None:
    """Create a real presentation with notes, a table, empty shape, picture, and chart."""

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide._element.cSld.set("name", "Executive Summary")
    slide.shapes.title.text = "Quarterly Overview"
    slide.notes_slide.notes_text_frame.text = "speaker notes 北区"

    table_shape = slide.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(4), Inches(1)
    )
    table_shape.table.cell(0, 0).text = "Region"
    table_shape.table.cell(0, 1).text = "Revenue"
    table_shape.table.cell(1, 0).text = "北区"
    table_shape.table.cell(1, 1).text = "120"

    empty_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(4), Inches(1), Inches(0.5)
    )
    empty_shape.name = "Empty semantic shape"

    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/"
        "fWQAAAAASUVORK5CYII="
    )
    slide.shapes.add_picture(
        _write_bytes(path.parent / "one-pixel.png", png),
        Inches(6),
        Inches(1),
        width=Inches(0.5),
        height=Inches(0.5),
    )

    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2"]
    chart_data.add_series("Sales", (3, 4))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(6),
        Inches(2),
        Inches(3),
        Inches(2),
        chart_data,
    )
    presentation.save(path)


def _inject_pptx_timing(path: Path) -> None:
    """Add valid unsupported timing/transition elements to the generated package."""

    temporary = path.with_suffix(".tmp.pptx")
    with (
        zipfile.ZipFile(path, "r") as source,
        zipfile.ZipFile(
            temporary, "w", compression=zipfile.ZIP_DEFLATED
        ) as destination,
    ):
        for info in source.infolist():
            contents = source.read(info.filename)
            if info.filename == "ppt/slides/slide1.xml":
                xml = contents.decode("utf-8")
                xml = xml.replace(
                    "</p:sld>",
                    "<p:transition/><p:timing/></p:sld>",
                )
                contents = xml.encode("utf-8")
            destination.writestr(info, contents)
    shutil.move(temporary, path)


def _write_xlsx(path: Path) -> None:
    """Create a real workbook with formula, table, merged, hidden, and view metadata."""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["Input", "Calculated"])
    worksheet.append([7, "=A2*3"])
    worksheet.append([10, "=SUM(A2:A3)"])
    worksheet.merge_cells("D1:E1")
    worksheet.row_dimensions[4].hidden = True
    worksheet.column_dimensions["F"].hidden = True
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.add_table(Table(displayName="ValuesTable", ref="A1:B3"))
    hidden = workbook.create_sheet("Hidden")
    hidden["A1"] = "not visible"
    hidden.sheet_state = "hidden"
    workbook.save(path)


def _set_cached_formula_values(
    path: Path,
    cached_values: dict[str, str],
    *,
    remove: set[str],
) -> None:
    """Write genuine OOXML formula caches for testing both cache branches."""

    temporary = path.with_suffix(".tmp.xlsx")
    with (
        zipfile.ZipFile(path, "r") as source,
        zipfile.ZipFile(
            temporary, "w", compression=zipfile.ZIP_DEFLATED
        ) as destination,
    ):
        for info in source.infolist():
            contents = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(contents)
                for cell in root.findall(".//main:c", _xlsx_ns()):
                    coordinate = cell.get("r")
                    if coordinate in remove:
                        value = cell.find("main:v", _xlsx_ns())
                        if value is not None:
                            cell.remove(value)
                    elif coordinate in cached_values:
                        value = cell.find("main:v", _xlsx_ns())
                        if value is None:
                            value = ElementTree.SubElement(cell, f"{{{SHEET_NS}}}v")
                        value.text = cached_values[coordinate]
                contents = ElementTree.tostring(
                    root, encoding="utf-8", xml_declaration=True
                )
            destination.writestr(info, contents)
    shutil.move(temporary, path)


def _xlsx_ns() -> dict[str, str]:
    """Return the namespace map used by the test's OOXML cache patcher."""

    return {"main": SHEET_NS}


def _write_bytes(path: Path, contents: bytes) -> str:
    """Persist one generated local image fixture and return its path."""

    path.write_bytes(contents)
    return str(path)
