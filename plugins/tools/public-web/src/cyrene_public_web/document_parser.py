"""
┌─────────────────────────────────────────────────────────────────────┐
│  📄 document_parser.py                                              │
│  Module: cyrene_public_web.document_parser                          │
│  Role: Extract bounded text blocks with licensed local libraries.   │
│                                                                     │
│  模块职责：使用许可明确的本地库提取有界文本块，不下载模型或伪造 OCR。  │
└─────────────────────────────────────────────────────────────────────┘
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import zipfile
from collections.abc import Sequence
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl
import pypdfium2
from docx import Document
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from PIL import Image, UnidentifiedImageError
from pptx import Presentation

MAX_SOURCE_BYTES = 10 * 1024 * 1024
MAX_OOXML_ENTRIES = 20_000
MAX_OOXML_EXPANDED_BYTES = 128 * 1024 * 1024
MAX_PDF_PAGES = 1_000
MAX_SLIDES = 1_000
MAX_SHEETS = 256
MAX_BLOCKS = 100_000
MAX_BLOCK_TEXT_CHARS = 4_096
MAX_BLOCKS_JSON_BYTES = 13 * 1024 * 1024
MAX_IMAGE_PIXELS = 100_000_000
_SHA256 = re.compile(r"sha256:[0-9a-f]{64}\Z")
_MARKDOWN_HEADING = re.compile(r"^ {0,3}(#{1,6})\s+(.+?)\s*#*\s*$")
_MEDIA_FORMATS = {
    "application/pdf": "PDF",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "DOCX",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "PPTX",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "XLSX",
    "text/csv": "CSV",
    "text/markdown": "MARKDOWN",
    "text/plain": "TXT",
    "image/png": "PNG",
    "image/jpeg": "JPEG",
}


class DocumentParseError(ValueError):
    """The staged source is invalid or exceeds a fixed parser limit."""


class _BlockBuilder:
    """Build stable, source-bound blocks and enforce the response budget."""

    def __init__(self, source_ref: str, source_digest: str) -> None:
        self.source_ref = source_ref
        self.source_digest = source_digest
        self.blocks: list[dict[str, Any]] = []
        self.warnings: list[str] = []
        self.diagnostics: list[dict[str, Any]] = []
        self.unsupported: list[dict[str, Any]] = []
        self._encoded_bytes = 0
        self._limited = False
        self._text_truncated = False

    @property
    def is_limited(self) -> bool:
        """Report when further source blocks would exceed the response budget."""

        return self._limited

    def add(
        self,
        text: str,
        *,
        kind: str,
        item_ref: str,
        source_pages: Sequence[int] = (),
        section_path: Sequence[str] = (),
        tree_level: int = 0,
        table_index: int | None = None,
    ) -> None:
        """Append one bounded block with a stable identity and public locator."""

        if self._limited or not text.strip():
            return
        if len(text) > MAX_BLOCK_TEXT_CHARS:
            text = text[:MAX_BLOCK_TEXT_CHARS]
            if not self._text_truncated:
                self.warn("One or more source blocks were truncated to the text limit.")
                self._text_truncated = True
        if len(self.blocks) >= MAX_BLOCKS:
            self._mark_limited()
            return
        ordinal = len(self.blocks)
        identity = hashlib.sha256(
            f"{self.source_ref}\0{self.source_digest}\0{ordinal}".encode()
        ).hexdigest()
        locator: dict[str, Any] = {
            "source_pages": list(source_pages[:100]),
            "section_path": [item[:256] for item in section_path[:32]],
            "item_ref": item_ref[:512],
            "tree_level": max(0, min(tree_level, 64)),
        }
        if table_index is not None and table_index >= 0:
            locator["table_index"] = table_index
        block = {
            "id": f"sha256:{identity}",
            "ordinal": ordinal,
            "kind": kind[:64],
            "text": text,
            "locator": locator,
        }
        encoded_size = len(
            json.dumps(block, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        )
        if self._encoded_bytes + encoded_size > MAX_BLOCKS_JSON_BYTES:
            self._mark_limited()
            return
        self.blocks.append(block)
        self._encoded_bytes += encoded_size

    def warn(self, message: str) -> None:
        if len(self.warnings) < 20:
            self.warnings.append(message[:1_000])

    def diagnostic(
        self,
        code: str,
        message: str,
        *,
        kind: str = "parser",
        source_page: int | None = None,
    ) -> None:
        if len(self.diagnostics) >= 100:
            return
        item: dict[str, Any] = {
            "code": code[:128],
            "message": message[:1_000],
            "kind": kind[:64],
        }
        if source_page is not None:
            item["source_pages"] = [source_page]
        self.diagnostics.append(item)

    def unsupported_content(
        self, kind: str, reason: str, *, source_page: int | None = None
    ) -> None:
        if len(self.unsupported) >= 100:
            return
        item: dict[str, Any] = {"kind": kind[:128], "reason": reason[:256]}
        if source_page is not None:
            item["source_pages"] = [source_page]
        self.unsupported.append(item)

    def _mark_limited(self) -> None:
        self._limited = True
        if "Block output limit reached; remaining source blocks were omitted." not in self.warnings:
            self.warn("Block output limit reached; remaining source blocks were omitted.")


def parse_document(
    *,
    source_path: Path,
    payload_path: Path,
    blocks_path: Path,
    result_path: Path,
    source_ref: str,
    source_digest: str,
    filename: str,
    media_type: str,
) -> dict[str, Any]:
    """Parse one staged document into bounded source blocks and a report.

    The worker receives only private staging paths and caller-owned source
    identity. Text extraction uses pinned, local libraries; OCR is reported as
    unavailable instead of returning synthetic text.
    中文：仅使用本地固定依赖读取文档；图片或无文本 PDF 不会伪造 OCR 结果。
    """

    source_path, output_paths = _validate_paths(source_path, payload_path, blocks_path, result_path)
    if not source_ref or "\n" in source_ref or "\r" in source_ref:
        raise DocumentParseError("source_ref is invalid")
    if not filename or len(filename) > 512 or "\n" in filename or "\r" in filename:
        raise DocumentParseError("filename is invalid")
    if not _SHA256.fullmatch(source_digest):
        raise DocumentParseError("source_digest must use sha256:<64 lowercase hex>")
    source_bytes = source_path.read_bytes()
    if not source_bytes or len(source_bytes) > MAX_SOURCE_BYTES:
        raise DocumentParseError("source exceeds its byte limit")
    if f"sha256:{hashlib.sha256(source_bytes).hexdigest()}" != source_digest:
        raise DocumentParseError("source digest does not match staged bytes")
    source_format = _MEDIA_FORMATS.get(media_type.split(";", maxsplit=1)[0].strip().lower())
    if source_format is None:
        raise DocumentParseError("media type is unsupported")

    builder = _BlockBuilder(source_ref, source_digest)
    summary: dict[str, Any] = {
        "source_format": source_format,
        "page_count": None,
        "slide_count": None,
        "sheet_count": None,
    }
    if source_format in {"TXT", "MARKDOWN"}:
        _parse_text(source_bytes, source_format, builder)
    elif source_format == "CSV":
        _parse_csv(source_bytes, builder)
    elif source_format == "PDF":
        summary["page_count"] = _parse_pdf(source_path, builder)
    elif source_format == "DOCX":
        _validate_ooxml(source_path)
        _parse_docx(source_path, builder)
    elif source_format == "PPTX":
        _validate_ooxml(source_path)
        summary["slide_count"] = _parse_pptx(source_path, builder)
    elif source_format == "XLSX":
        _validate_ooxml(source_path)
        summary["sheet_count"] = _parse_xlsx(source_path, builder)
    else:
        summary["page_count"] = 1
        _inspect_image(source_path, source_format, builder)

    block_envelope = {
        "schema_version": "cyrene.public-web.document-blocks.v1",
        "source_ref": source_ref,
        "source_digest": source_digest,
        "blocks": builder.blocks,
    }
    report = {
        "source_ref": source_ref,
        "source_digest": source_digest,
        "warnings": builder.warnings,
        "diagnostics": builder.diagnostics,
        "unsupported_content": builder.unsupported,
    }
    _write_json(output_paths[0], {"schema_version": "cyrene.public-web.document-payload.v1"})
    _write_json(output_paths[1], block_envelope)
    _write_json(output_paths[2], report)
    return summary


def _validate_paths(
    source_path: Path,
    payload_path: Path,
    blocks_path: Path,
    result_path: Path,
) -> tuple[Path, tuple[Path, Path, Path]]:
    """Require one regular input and fresh outputs inside one staging directory."""

    source = Path(source_path)
    if source.is_symlink() or not source.is_file():
        raise DocumentParseError("source_path must point to a regular file")
    parent = source.parent.resolve(strict=True)
    if not source.resolve(strict=True).is_relative_to(parent):
        raise DocumentParseError("source path escapes its private staging directory")
    if source.stat().st_size > MAX_SOURCE_BYTES:
        raise DocumentParseError("source exceeds its byte limit")
    outputs = tuple(Path(path) for path in (payload_path, blocks_path, result_path))
    for path in outputs:
        if not path.is_absolute() or path.parent.resolve() != parent:
            raise DocumentParseError("output path is outside private staging")
        if path.exists() or path.is_symlink():
            raise DocumentParseError("output path already exists or is unsafe")
    return source, outputs  # type: ignore[return-value]


def _write_json(path: Path, value: dict[str, Any]) -> None:
    """Create one private, bounded JSON file without following a symlink."""

    payload = json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":")).encode(
        "utf-8"
    )
    if len(payload) > MAX_BLOCKS_JSON_BYTES:
        raise DocumentParseError("parser output exceeds its byte limit")
    descriptor = os.open(
        path,
        os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, "O_NOFOLLOW", 0),
        0o600,
    )
    with os.fdopen(descriptor, "wb") as stream:
        stream.write(payload)


def _decode_text(payload: bytes) -> str:
    try:
        return payload.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as error:
        raise DocumentParseError("text source is not valid UTF-8") from error


def _parse_text(payload: bytes, source_format: str, builder: _BlockBuilder) -> None:
    text = _decode_text(payload)
    headings: list[str] = []
    for line_number, line in enumerate(text.splitlines(), start=1):
        if source_format == "MARKDOWN":
            heading = _MARKDOWN_HEADING.fullmatch(line)
            if heading is not None:
                level = len(heading.group(1))
                headings = headings[: level - 1]
                headings.append(heading.group(2).strip())
                builder.add(
                    line,
                    kind="heading",
                    item_ref=f"line:{line_number}",
                    section_path=headings,
                    tree_level=level,
                )
                if builder.is_limited:
                    break
                continue
        builder.add(
            line,
            kind="line",
            item_ref=f"line:{line_number}",
            section_path=headings,
            tree_level=len(headings),
        )
        if builder.is_limited:
            break


def _parse_csv(payload: bytes, builder: _BlockBuilder) -> None:
    text = _decode_text(payload)
    try:
        for row_number, row in enumerate(csv.reader(io.StringIO(text, newline="")), start=1):
            builder.add(
                " | ".join(row),
                kind="record",
                item_ref=f"row:{row_number}",
                table_index=0,
            )
            if builder.is_limited:
                break
    except csv.Error as error:
        raise DocumentParseError("CSV source is malformed") from error


def _parse_pdf(source_path: Path, builder: _BlockBuilder) -> int:
    document = pypdfium2.PdfDocument(str(source_path))
    try:
        page_count = len(document)
        if page_count > MAX_PDF_PAGES:
            builder.warn("PDF page limit reached; remaining pages were not read.")
        for page_index in range(min(page_count, MAX_PDF_PAGES)):
            page = document.get_page(page_index)
            try:
                text_page = page.get_textpage()
                try:
                    text = text_page.get_text_range()
                finally:
                    text_page.close()
            finally:
                page.close()
            page_number = page_index + 1
            lines = [line for line in text.splitlines() if line.strip()]
            if not lines:
                builder.diagnostic(
                    "OCR_NOT_CONFIGURED",
                    "No searchable text was found; image text extraction was not run.",
                    kind="ocr",
                    source_page=page_number,
                )
                builder.unsupported_content(
                    "image_text",
                    "ocr_not_configured",
                    source_page=page_number,
                )
                continue
            for line_number, line in enumerate(lines, start=1):
                builder.add(
                    line,
                    kind="line",
                    item_ref=f"page:{page_number}:line:{line_number}",
                    source_pages=(page_number,),
                )
                if builder.is_limited:
                    break
            if builder.is_limited:
                break
        return page_count
    finally:
        document.close()


def _validate_ooxml(source_path: Path) -> None:
    """Bound the declared ZIP expansion before passing Office input to a library."""

    try:
        with zipfile.ZipFile(source_path) as archive:
            infos = archive.infolist()
    except (OSError, zipfile.BadZipFile) as error:
        raise DocumentParseError("Office source is not a valid ZIP package") from error
    if len(infos) > MAX_OOXML_ENTRIES:
        raise DocumentParseError("Office source contains too many ZIP entries")
    if sum(info.file_size for info in infos) > MAX_OOXML_EXPANDED_BYTES:
        raise DocumentParseError("Office source expanded size exceeds its limit")
    if any(info.flag_bits & 0x1 for info in infos):
        raise DocumentParseError("encrypted Office archives are unsupported")


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = getattr(paragraph.style, "name", "") or ""
    match = re.fullmatch(r"Heading ([1-6])", style_name)
    return None if match is None else int(match.group(1))


def _parse_docx(source_path: Path, builder: _BlockBuilder) -> None:
    document = Document(str(source_path))
    headings: list[str] = []
    table_index = 0
    paragraph_index = 0
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            paragraph_index += 1
            text = item.text
            level = _heading_level(item)
            if level is not None and text.strip():
                headings = headings[: level - 1]
                headings.append(text.strip())
                builder.add(
                    text,
                    kind="heading",
                    item_ref=f"paragraph:{paragraph_index}",
                    section_path=headings,
                    tree_level=level,
                )
                if builder.is_limited:
                    return
            else:
                builder.add(
                    text,
                    kind="paragraph",
                    item_ref=f"paragraph:{paragraph_index}",
                    section_path=headings,
                    tree_level=len(headings),
                )
                if builder.is_limited:
                    return
        elif isinstance(item, DocxTable):
            for row_number, row in enumerate(item.rows, start=1):
                for cell_number, cell in enumerate(row.cells, start=1):
                    builder.add(
                        cell.text,
                        kind="table_cell",
                        item_ref=f"table:{table_index}:row:{row_number}:cell:{cell_number}",
                        section_path=headings,
                        tree_level=len(headings),
                        table_index=table_index,
                    )
                    if builder.is_limited:
                        return
            table_index += 1


def _parse_pptx(source_path: Path, builder: _BlockBuilder) -> int:
    presentation = Presentation(str(source_path))
    slide_count = len(presentation.slides)
    if slide_count > MAX_SLIDES:
        builder.warn("Presentation slide limit reached; remaining slides were not read.")
    for slide_index in range(min(slide_count, MAX_SLIDES)):
        slide = presentation.slides[slide_index]
        slide_number = slide_index + 1
        for shape_index, shape in enumerate(slide.shapes, start=1):
            if shape.has_text_frame:
                text = shape.text
                if text.strip():
                    builder.add(
                        text,
                        kind="slide_text",
                        item_ref=f"slide:{slide_number}:shape:{shape_index}",
                        source_pages=(slide_number,),
                    )
                    if builder.is_limited:
                        return slide_count
            if shape.has_table:
                for row_index, row in enumerate(shape.table.rows, start=1):
                    for cell_index, cell in enumerate(row.cells, start=1):
                        builder.add(
                            cell.text,
                            kind="table_cell",
                            item_ref=f"slide:{slide_number}:table:{shape_index}:row:{row_index}:cell:{cell_index}",
                            source_pages=(slide_number,),
                            table_index=shape_index - 1,
                        )
                        if builder.is_limited:
                            return slide_count
    return slide_count


def _cell_text(value: object) -> str:
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


def _parse_xlsx(source_path: Path, builder: _BlockBuilder) -> int:
    workbook = openpyxl.load_workbook(
        source_path, read_only=True, data_only=False, keep_links=False
    )
    try:
        sheet_count = len(workbook.worksheets)
        if sheet_count > MAX_SHEETS:
            builder.warn("Spreadsheet sheet limit reached; remaining sheets were not read.")
        for sheet_index, worksheet in enumerate(workbook.worksheets[:MAX_SHEETS]):
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    text = _cell_text(cell.value)
                    builder.add(
                        text,
                        kind="cell",
                        item_ref=f"sheet:{worksheet.title}:cell:{cell.coordinate}",
                        section_path=(worksheet.title,),
                        table_index=sheet_index,
                    )
                    if builder.is_limited:
                        return sheet_count
        return sheet_count
    finally:
        workbook.close()


def _inspect_image(source_path: Path, source_format: str, builder: _BlockBuilder) -> None:
    try:
        with Image.open(source_path) as image:
            expected_format = "JPEG" if source_format == "JPEG" else "PNG"
            if image.format != expected_format:
                raise DocumentParseError("image bytes do not match their media type")
            width, height = image.size
            if width < 1 or height < 1 or width * height > MAX_IMAGE_PIXELS:
                raise DocumentParseError("image dimensions exceed the parser limit")
            image.verify()
    except (OSError, UnidentifiedImageError) as error:
        raise DocumentParseError("image source is invalid") from error
    builder.warn("Image text was not extracted because OCR is not configured.")
    builder.diagnostic(
        "OCR_NOT_CONFIGURED",
        "Image dimensions were verified; OCR was not run.",
        kind="ocr",
        source_page=1,
    )
    builder.unsupported_content("image_text", "ocr_not_configured", source_page=1)
