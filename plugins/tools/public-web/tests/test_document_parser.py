"""
┌─────────────────────────────────────────────────────────────────────┐
│  📄 test_document_parser.py                                          │
│  Module: plugins.tools.public_web.tests                              │
│  Role: Verify the isolated local document extraction formats.        │
│                                                                     │
│  模块职责：验证隔离本地解析器的文本、PDF、Office 与无 OCR 图片边界。  │
└─────────────────────────────────────────────────────────────────────┘
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import subprocess
import sys
from pathlib import Path
from typing import Any

import openpyxl
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from cyrene_public_web import document_parser

WORKER = Path(document_parser.__file__).with_name("document_parser_worker.py")
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _minimal_pdf(text: str) -> bytes:
    """Create one small searchable PDF page without a PDF authoring dependency."""

    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(content)).encode("ascii")
        + b" >>\nstream\n"
        + content
        + b"\nendstream",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode("ascii"))
        payload.extend(body)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(payload)


def _worker_environment(tmp_path: Path) -> dict[str, str]:
    environment = {
        "PATH": os.environ.get("PATH", ""),
        "HOME": str(tmp_path),
        "PYTHONNOUSERSITE": "1",
        "PYTHONUTF8": "1",
        "HF_HUB_OFFLINE": "1",
        "TRANSFORMERS_OFFLINE": "1",
    }
    for name in ("SYSTEMROOT", "WINDIR", "TEMP", "TMP"):
        if name in os.environ:
            environment[name] = os.environ[name]
    return environment


def _parse_in_worker(
    tmp_path: Path, *, suffix: str, media_type: str, source: bytes
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    tmp_path.mkdir(parents=True, exist_ok=True)
    source_path = tmp_path / f"source{suffix}"
    source_path.write_bytes(source)
    source_digest = f"sha256:{hashlib.sha256(source).hexdigest()}"
    source_ref = f"sha256:{hashlib.sha256(('source:' + source_digest).encode()).hexdigest()}"
    request_path = tmp_path / "request.json"
    summary_path = tmp_path / "summary.json"
    blocks_path = tmp_path / "blocks.json"
    result_path = tmp_path / "result.json"
    request_path.write_text(
        json.dumps(
            {
                "source_path": str(source_path),
                "payload_path": str(tmp_path / "payload.json"),
                "blocks_path": str(blocks_path),
                "result_path": str(result_path),
                "source_ref": source_ref,
                "source_digest": source_digest,
                "filename": source_path.name,
                "media_type": media_type,
            },
            separators=(",", ":"),
        ),
        encoding="utf-8",
    )
    environment = _worker_environment(tmp_path)
    assert "PYTHONPATH" not in environment
    process = subprocess.run(
        [sys.executable, "-B", str(WORKER), str(request_path)],
        cwd=tmp_path,
        env=environment,
        capture_output=True,
        text=True,
        check=False,
        timeout=30,
    )
    assert process.returncode == 0, process.stderr
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    blocks = json.loads(blocks_path.read_text(encoding="utf-8"))
    report = json.loads(result_path.read_text(encoding="utf-8"))
    return summary, blocks, report


def test_text_and_csv_extract_source_bound_blocks(tmp_path: Path) -> None:
    summary, blocks, report = _parse_in_worker(
        tmp_path,
        suffix=".txt",
        media_type="text/plain",
        source=b"first line\nsecond line\n",
    )
    assert summary["source_format"] == "TXT"
    assert blocks["schema_version"] == "cyrene.public-web.document-blocks.v1"
    assert [block["text"] for block in blocks["blocks"]] == [
        "first line",
        "second line",
    ]
    assert blocks["blocks"][0]["locator"]["item_ref"] == "line:1"
    assert blocks["source_digest"] == report["source_digest"]

    csv_root = tmp_path / "csv"
    csv_root.mkdir()
    _, csv_blocks, _ = _parse_in_worker(
        csv_root,
        suffix=".csv",
        media_type="text/csv",
        source=b"name,value\nAda,1\n",
    )
    assert [block["text"] for block in csv_blocks["blocks"]] == [
        "name | value",
        "Ada | 1",
    ]
    assert csv_blocks["blocks"][1]["locator"]["item_ref"] == "row:2"


def test_markdown_headings_are_preserved_as_source_sections(tmp_path: Path) -> None:
    summary, blocks, _ = _parse_in_worker(
        tmp_path,
        suffix=".md",
        media_type="text/markdown",
        source=b"# Overview\nA bounded paragraph.\n## Details\nMore text.\n",
    )
    assert summary["source_format"] == "MARKDOWN"
    assert blocks["blocks"][0]["kind"] == "heading"
    assert blocks["blocks"][1]["locator"]["section_path"] == ["Overview"]
    assert blocks["blocks"][3]["locator"]["section_path"] == ["Overview", "Details"]


def test_pdf_text_layer_is_read_without_ocr_or_model_download(tmp_path: Path) -> None:
    summary, blocks, report = _parse_in_worker(
        tmp_path,
        suffix=".pdf",
        media_type="application/pdf",
        source=_minimal_pdf("Searchable PDF text"),
    )
    assert summary["source_format"] == "PDF"
    assert summary["page_count"] == 1
    assert any("Searchable PDF text" in block["text"] for block in blocks["blocks"])
    assert blocks["blocks"][0]["locator"]["source_pages"] == [1]
    assert report["diagnostics"] == []


def test_docx_pptx_and_xlsx_extract_real_office_text(tmp_path: Path) -> None:
    docx_path = tmp_path / "source.docx"
    document = Document()
    document.add_heading("Project Notes", level=1)
    document.add_paragraph("A local paragraph.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "A table cell."
    document.save(docx_path)
    _, docx_blocks, _ = _parse_in_worker(
        tmp_path / "docx-run",
        suffix=".docx",
        media_type=DOCX_MEDIA_TYPE,
        source=docx_path.read_bytes(),
    )
    docx_text = [block["text"] for block in docx_blocks["blocks"]]
    assert "Project Notes" in docx_text
    assert "A local paragraph." in docx_text
    assert "A table cell." in docx_text
    assert docx_blocks["blocks"][1]["locator"]["section_path"] == ["Project Notes"]

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Slide text"
    pptx_path = tmp_path / "source.pptx"
    presentation.save(pptx_path)
    pptx_root = tmp_path / "pptx-run"
    pptx_root.mkdir()
    pptx_summary, pptx_blocks, _ = _parse_in_worker(
        pptx_root,
        suffix=".pptx",
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        source=pptx_path.read_bytes(),
    )
    assert pptx_summary["slide_count"] == 1
    assert pptx_blocks["blocks"][0]["text"] == "Slide text"
    assert pptx_blocks["blocks"][0]["locator"]["source_pages"] == [1]

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Scores"
    worksheet["A1"] = "Player"
    worksheet["A2"] = "Ada"
    worksheet["B2"] = 12
    worksheet["C2"] = "=B2*2"
    xlsx_path = tmp_path / "source.xlsx"
    workbook.save(xlsx_path)
    xlsx_root = tmp_path / "xlsx-run"
    xlsx_root.mkdir()
    xlsx_summary, xlsx_blocks, _ = _parse_in_worker(
        xlsx_root,
        suffix=".xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        source=xlsx_path.read_bytes(),
    )
    assert xlsx_summary["sheet_count"] == 1
    assert [block["text"] for block in xlsx_blocks["blocks"]] == [
        "Player",
        "Ada",
        "12",
        "=B2*2",
    ]
    assert xlsx_blocks["blocks"][1]["locator"]["item_ref"] == "sheet:Scores:cell:A2"


def test_images_are_validated_without_fabricated_ocr(tmp_path: Path) -> None:
    image_bytes = io.BytesIO()
    Image.new("RGB", (8, 8), color="white").save(image_bytes, format="PNG")
    summary, blocks, report = _parse_in_worker(
        tmp_path,
        suffix=".png",
        media_type="image/png",
        source=image_bytes.getvalue(),
    )
    assert summary["source_format"] == "PNG"
    assert blocks["blocks"] == []
    assert report["diagnostics"][0]["code"] == "OCR_NOT_CONFIGURED"
    assert report["unsupported_content"][0]["reason"] == "ocr_not_configured"


def test_invalid_office_archive_fails_without_a_fake_result(tmp_path: Path) -> None:
    source_path = tmp_path / "source.docx"
    source_path.write_bytes(b"not a zip archive")
    digest = f"sha256:{hashlib.sha256(source_path.read_bytes()).hexdigest()}"
    request = tmp_path / "request.json"
    request.write_text(
        json.dumps(
            {
                "source_path": str(source_path),
                "payload_path": str(tmp_path / "payload.json"),
                "blocks_path": str(tmp_path / "blocks.json"),
                "result_path": str(tmp_path / "result.json"),
                "source_ref": "package-test-source",
                "source_digest": digest,
                "filename": source_path.name,
                "media_type": DOCX_MEDIA_TYPE,
            }
        ),
        encoding="utf-8",
    )
    result = subprocess.run(
        [sys.executable, "-B", str(WORKER), str(request)],
        cwd=tmp_path,
        env=_worker_environment(tmp_path),
        capture_output=True,
        check=False,
        timeout=10,
    )
    assert result.returncode == 1
    assert not (tmp_path / "blocks.json").exists()
    assert not (tmp_path / "result.json").exists()
